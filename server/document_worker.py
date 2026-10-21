#!/usr/bin/env python3
"""Isolated, stateless document transformations. JSON in/out; no user files written to disk."""
import base64
import io
import json
import re
import sys
import zipfile

MAX_FILE = 6 * 1024 * 1024
FORMATS = {"docx", "pptx", "xlsx", "pdf"}

def require(condition, message):
    if not condition:
        raise ValueError(message)

def text(value, limit=4000):
    require(isinstance(value, str) and len(value) <= limit, "Invalid text value")
    return value

def zip_check(data, kind):
    require(zipfile.is_zipfile(io.BytesIO(data)), "Invalid Office document archive")
    prefix = {"docx": "word/document.xml", "pptx": "ppt/presentation.xml", "xlsx": "xl/workbook.xml"}[kind]
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        infos = archive.infolist()
        require(len(infos) <= 1200, "Too many archive entries")
        require(sum(e.file_size for e in infos) <= 40 * 1024 * 1024, "Document expands beyond safety limit")
        require(max((e.file_size for e in infos), default=0) <= 16 * 1024 * 1024, "Oversized embedded document part")
        require(prefix in archive.namelist(), "Document type does not match its extension")

def preview_docx(data):
    from docx import Document
    d = Document(io.BytesIO(data))
    require(len(d.paragraphs) <= 1500 and len(d.tables) <= 100, "Document exceeds editing limits")
    return {"paragraphs": [{"index": i, "text": p.text[:4000]} for i,p in enumerate(d.paragraphs[:300])],
        "paragraphCount": len(d.paragraphs),
        "tables": [{"index": t, "rows": [[cell.text[:800] for cell in row.cells[:24]] for row in table.rows[:60]]}
            for t,table in enumerate(d.tables[:20])],
        "notes": ["Edits to a paragraph may alter its inline run formatting. Other content is preserved where supported."]}

def create_docx(model):
    from docx import Document
    doc=Document()
    if "markdown" in model:
        markdown=text(model["markdown"],12000)
        lines=markdown.splitlines()
        require(len(lines)<=250 and all(len(line)<=1600 for line in lines),
                "Markdown document exceeds line limits")
        for line in lines:
            heading=re.match(r"^(#{1,4})\s+(.+)$",line)
            if heading:
                doc.add_heading(heading.group(2).strip(),level=min(len(heading.group(1)),3))
                continue
            bullet=re.match(r"^\s*[-*]\s+(.+)$",line)
            ordered=re.match(r"^\s*[0-9]+[.]\s+(.+)$",line)
            label=bullet.group(1) if bullet else ordered.group(1) if ordered else line
            paragraph=doc.add_paragraph(style="List Bullet" if bullet else "List Number" if ordered else None)
            for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\x60[^\x60]+\x60)",label):
                if not part: continue
                run=paragraph.add_run(part[2:-2] if part.startswith("**") and part.endswith("**") else
                                      part[1:-1] if part.startswith(("*","\x60")) and part.endswith(("*","\x60")) else part)
                if part.startswith("**") and part.endswith("**"):run.bold=True
                elif part.startswith("*") and part.endswith("*"):run.italic=True
        out=io.BytesIO();doc.save(out);return out.getvalue()
    for p in model.get("paragraphs", [])[:300]:
        doc.add_paragraph(text(p,4000))
    for table in model.get("tables", [])[:20]:
        rows=table.get("rows", [])[:60]
        if not rows: continue
        width=min(max(len(r) for r in rows),24)
        dtable=doc.add_table(rows=len(rows),cols=width)
        for i,row in enumerate(rows):
            for j,value in enumerate(row[:width]):dtable.cell(i,j).text=text(str(value),800)
    out=io.BytesIO();doc.save(out);return out.getvalue()

def replace_styled_runs(paragraph, replacement):
    old = paragraph.text
    if old == replacement: return
    runs = list(paragraph.runs)
    if not runs:
        require(not old, "Unsupported inline content")
        paragraph.add_run(replacement); return
    require("".join(run.text for run in runs) == old,
            "Unsupported inline content; cannot safely edit")
    prefix = 0
    while prefix < min(len(old),len(replacement)) and old[prefix] == replacement[prefix]:
        prefix += 1
    suffix = 0
    while (suffix < len(old)-prefix and suffix < len(replacement)-prefix
           and old[-1-suffix] == replacement[-1-suffix]):
        suffix += 1
    suffix_start = len(old)-suffix
    inserted = replacement[prefix:len(replacement)-suffix if suffix else len(replacement)]
    offset = 0
    spans = []
    for run in runs:
        spans.append((offset,offset+len(run.text)))
        offset += len(run.text)
    anchor = next((i for i,(a,b) in enumerate(spans) if a <= prefix < b),len(runs)-1)
    updates = []
    allowed = {"rPr","t","br","cr","tab"}
    for i,(run,(start,end)) in enumerate(zip(runs,spans)):
        original = run.text
        left = original[:max(0,min(len(original),prefix-start))]
        right = original[max(0,min(len(original),suffix_start-start)):]
        revised = left + (inserted if i==anchor else "") + right
        if revised != original:
            require(all(child.tag.rsplit("}",1)[-1] in allowed for child in run._r),
                    "Inline drawing or field prevents safe editing")
            updates.append((run,revised))
    for run,revised in updates: run.text = revised

def edit_docx(data, changes):
    from docx import Document
    doc=Document(io.BytesIO(data))
    for edit in changes:
        value=text(edit.get("text"),4000)
        if edit.get("target")=="paragraph":
            i=edit.get("index")
            require(type(i)==int and 0<=i<len(doc.paragraphs),"Paragraph index out of range")
            replace_styled_runs(doc.paragraphs[i],value)
        elif edit.get("target")=="cell":
            t,r,c=(edit.get(k) for k in ("table","row","col"))
            require(all(type(x)==int and x>=0 for x in (t,r,c)), "Invalid table cell location")
            require(t<len(doc.tables) and r<len(doc.tables[t].rows) and c<len(doc.tables[t].rows[r].cells),"Cell outside table")
            cell=doc.tables[t].cell(r,c)
            require(len(cell.paragraphs)==1,"Editing multi-paragraph table cells is not supported")
            replace_styled_runs(cell.paragraphs[0],value)
        else: raise ValueError("Unsupported DOCX edit")
    out=io.BytesIO();doc.save(out);return out.getvalue()

def preview_pptx(data):
    from pptx import Presentation
    prs=Presentation(io.BytesIO(data))
    require(len(prs.slides)<=250,"Too many slides")
    return {"slides":[{"index":i,"shapes":[{"index":j,"text":shape.text[:4000]}
        for j,shape in enumerate(slide.shapes) if shape.has_text_frame][:30]}
        for i,slide in enumerate(prs.slides) if i<60],
        "slideCount":len(prs.slides),
        "notes":["Editing shape text can change its text styling. Non-text shapes are retained."]}

def create_pptx(model):
    from pptx import Presentation
    prs=Presentation()
    slides=model.get("slides", [])
    require(isinstance(slides,list) and len(slides)<=60,"Too many slides")
    for slide in slides:
        require(isinstance(slide,dict),"Invalid slide")
        s=prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text=text(slide.get("title",""),240)
        s.placeholders[1].text=text(slide.get("body",""),4000)
    out=io.BytesIO();prs.save(out);return out.getvalue()

def edit_pptx(data,changes):
    from pptx import Presentation
    prs=Presentation(io.BytesIO(data))
    for edit in changes:
        i,j=edit.get("slide"),edit.get("shape")
        require(type(i)==int and type(j)==int and 0<=i<len(prs.slides) and 0<=j<len(prs.slides[i].shapes),"Invalid shape location")
        shape=prs.slides[i].shapes[j]
        require(shape.has_text_frame,"Shape has no editable text")
        replacement=text(edit.get("text"),4000)
        pieces=replacement.split("\n")
        paragraphs=list(shape.text_frame.paragraphs)
        require(len(pieces)==len(paragraphs),"Changing slide paragraph count can alter styling")
        for paragraph,piece in zip(paragraphs,pieces):
            replace_styled_runs(paragraph,piece)
    out=io.BytesIO();prs.save(out);return out.getvalue()

def value_repr(value):
    if value is None:return ""
    if isinstance(value,(str,int,float,bool)):return str(value)[:1200]
    return str(value)[:1200]

def preview_xlsx(data):
    import openpyxl
    wb=openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
    require(len(wb.worksheets)<=50,"Too many worksheets")
    sheets=[]
    for ws in wb.worksheets[:15]:
        cells=[]
        for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,60),min_col=1,max_col=min(ws.max_column,24)):
            for cell in row:
                if cell.value is not None:cells.append({"address":cell.coordinate,"value":value_repr(cell.value)})
                if len(cells)>=1000:break
            if len(cells)>=1000:break
        sheets.append({"name":ws.title,"cells":cells,"rows":ws.max_row,"columns":ws.max_column})
    return {"sheets":sheets,"sheetCount":len(wb.worksheets),
        "notes":["Only a bounded cell preview is shown; other cells remain in the file. Formula results are not recalculated."]}

def create_xlsx(model):
    import openpyxl
    wb=openpyxl.Workbook()
    sheets=model.get("sheets",[])
    require(isinstance(sheets,list) and 1<=len(sheets)<=15,"Provide 1–15 sheets")
    for i,sheet in enumerate(sheets):
        ws=wb.active if i==0 else wb.create_sheet()
        ws.title=text(sheet.get("name",f"Sheet{i+1}"),30) or f"Sheet{i+1}"
        cells=sheet.get("cells",[])
        require(isinstance(cells,list) and len(cells)<=1000,"Too many cells")
        for c in cells:
            ws[c["address"]]=text(c.get("value",""),1200)
    out=io.BytesIO();wb.save(out);return out.getvalue()

def edit_xlsx(data,changes):
    import openpyxl
    from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
    wb=openpyxl.load_workbook(io.BytesIO(data), read_only=False, keep_links=False)
    for edit in changes:
        name=edit.get("sheet")
        require(isinstance(name,str) and name in wb.sheetnames,"Unknown worksheet")
        address=edit.get("address")
        require(isinstance(address,str) and re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,5}",address)!=None,"Invalid cell address")
        col,row=coordinate_from_string(address)
        require(column_index_from_string(col)<=1000 and row<=100000,"Cell outside editing limits")
        wb[name][address]=text(edit.get("value",""),1200)
    out=io.BytesIO();wb.save(out);return out.getvalue()

def pdf_create(lines):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.enums import TA_LEFT
    from xml.sax.saxutils import escape
    buf=io.BytesIO()
    doc=SimpleDocTemplate(buf, pageCompression=1, title="GoreeCloud document")
    style=getSampleStyleSheet()["BodyText"]
    parts=[]
    for paragraph in lines.splitlines()[:250]:
        parts.append(Paragraph(escape(paragraph[:1600]) or " ",style))
        parts.append(Spacer(1,6))
    if not parts:parts=[Paragraph(" ",style)]
    doc.build(parts)
    return buf.getvalue()

def preview_pdf(data):
    from pypdf import PdfReader
    pdf=PdfReader(io.BytesIO(data),strict=True)
    require(not pdf.is_encrypted,"Encrypted PDF requires local decryption before editing")
    require(len(pdf.pages)<=250,"PDF exceeds page limits")
    return {"pages":[{"index":i,"rotation":page.rotation,"text":(page.extract_text() or "")[:5000]}
      for i,page in enumerate(pdf.pages[:35])],
      "pageCount":len(pdf.pages),"notes":["Existing PDF text is preview-only. Editing supports appending text pages or page rotation, not in-place rewriting."]}

def edit_pdf(data,changes):
    from pypdf import PdfReader,PdfWriter
    reader=PdfReader(io.BytesIO(data),strict=True)
    require(not reader.is_encrypted,"Encrypted PDF requires decryption")
    require(len(reader.pages)<=250,"Too many PDF pages")
    writer=PdfWriter()
    for page in reader.pages:writer.add_page(page)
    for edit in changes:
        action=edit.get("target")
        if action=="append":
            value=text(edit.get("text"),12000)
            extra=PdfReader(io.BytesIO(pdf_create(value)))
            for page in extra.pages:writer.add_page(page)
        elif action=="rotate":
            i=edit.get("index")
            degrees=edit.get("degrees")
            require(type(i)==int and 0<=i<len(writer.pages) and degrees in (90,180,270),"Invalid PDF rotation")
            writer.pages[i].rotate(degrees)
        else:raise ValueError("Unsupported PDF change")
    out=io.BytesIO();writer.write(out);return out.getvalue()

def transform(request):
    action=request.get("action")
    name=request.get("name")
    require(isinstance(name,str) and re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9 _.()\-]{0,100}\.[A-Za-z0-9]{2,5}",name)!=None,"Invalid filename")
    kind=name.rsplit(".",1)[1].lower()
    require(kind in FORMATS,"Unsupported document type")
    require(action in ("open","create","edit"),"Unsupported document operation")
    data=b""
    if action in ("open","edit"):
        encoded=request.get("base64","")
        require(isinstance(encoded,str) and len(encoded)<=MAX_FILE*4//3+8,"Document upload too large")
        data=base64.b64decode(encoded,validate=True)
        require(0<len(data)<=MAX_FILE,"Document upload too large or empty")
        if kind=="pdf":require(data.startswith(b"%PDF-"),"File is not a PDF")
        else:zip_check(data,kind)
    if action=="open":
        model={"docx":preview_docx,"pptx":preview_pptx,"xlsx":preview_xlsx,"pdf":preview_pdf}[kind](data)
        return {"kind":kind,"model":model}
    if action=="create":
        model=request.get("model",{})
        require(isinstance(model,dict),"Invalid document model")
        if kind=="pdf":
            if model.get("markdown") is True:
                from importlib.util import spec_from_file_location, module_from_spec
                from pathlib import Path
                _spec=spec_from_file_location('goreecloud_markdown_pdf',Path(__file__).with_name('markdown_pdf.py'))
                _module=module_from_spec(_spec)
                _spec.loader.exec_module(_module)
                markdown_to_pdf=_module.markdown_to_pdf
                result=markdown_to_pdf(text(model.get("text",""),12000))
            else:result=pdf_create(text(model.get("text",""),12000))
        else:result={"docx":create_docx,"pptx":create_pptx,"xlsx":create_xlsx}[kind](model)
    else:
        changes=request.get("changes",[])
        require(isinstance(changes,list) and 1<=len(changes)<=120,"Provide 1–120 changes")
        result={"docx":edit_docx,"pptx":edit_pptx,"xlsx":edit_xlsx,"pdf":edit_pdf}[kind](data,changes)
    require(len(result)<=MAX_FILE,"Generated file exceeds size limit")
    return {"kind":kind,"base64":base64.b64encode(result).decode("ascii"),"bytes":len(result)}

def main():
    try:
        request=json.load(sys.stdin)
        require(isinstance(request,dict),"Invalid request")
        result=transform(request)
        json.dump({"ok":True,**result},sys.stdout)
    except Exception as e:
        json.dump({"ok":False,"error":str(e)[:220]},sys.stdout)
    sys.stdout.write("\n")

if __name__=="__main__":main()
