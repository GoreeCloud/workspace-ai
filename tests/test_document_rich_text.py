#!/usr/bin/env python3
import io,sys,base64
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'server'))
from document_worker import transform
from docx import Document
from pptx import Presentation
def modify(kind,data,changes):
 result=transform({'action':'edit','name':'styled.'+kind,
  'base64':base64.b64encode(data).decode(),'changes':changes})
 return base64.b64decode(result['base64'])

doc=Document()
paragraph=doc.add_paragraph()
paragraph.add_run('Bold ').bold=True
paragraph.add_run('word').italic=True
buffer=io.BytesIO();doc.save(buffer)
out=modify('docx',buffer.getvalue(),[{'target':'paragraph','index':0,'text':'Bold earth'}])
runs=Document(io.BytesIO(out)).paragraphs[0].runs
assert runs[0].text=='Bold ' and runs[0].bold is True
assert runs[1].text=='earth' and runs[1].italic is True
from pptx.util import Inches
presentation=Presentation()
slide=presentation.slides.add_slide(presentation.slide_layouts[6])
shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(6),Inches(2))
paragraph=shape.text_frame.paragraphs[0]
run=paragraph.add_run();run.text='Bold ';run.font.bold=True
run=paragraph.add_run();run.text='word';run.font.italic=True
buffer=io.BytesIO();presentation.save(buffer)
out=modify('pptx',buffer.getvalue(),[{'target':'shape','slide':0,'shape':0,'text':'Bold earth'}])
runs=Presentation(io.BytesIO(out)).slides[0].shapes[0].text_frame.paragraphs[0].runs
assert runs[0].text=='Bold ' and runs[0].font.bold is True
assert runs[1].text=='earth' and runs[1].font.italic is True
try:
    modify('pptx',out,[{'target':'shape','slide':0,'shape':0,'text':'First'+chr(10)+'Second'}])
    raise AssertionError('Expected multi-paragraph edit to be rejected')
except ValueError as error:
    assert 'paragraph count' in str(error)
print('RICH_TEXT_ROUNDTRIP_PASS')
